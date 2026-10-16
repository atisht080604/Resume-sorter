import re
import os
from typing import Dict, Any, List, Tuple

class DocumentParsingError(Exception):
    pass

class DocumentParser:
    """
    Extracts and normalizes text from PDF and DOCX documents.
    Detects sections, contact info, education, and experience indicators.
    """

    SECTION_PATTERNS = {
        'skills': [
            r'\b(?:technical\s+skills|core\s+competencies|key\s+skills|skills\s*(?:&|and)\s*tools|skills|technologies)\b'
        ],
        'experience': [
            r'\b(?:work\s+experience|professional\s+experience|employment\s+history|work\s+history|experience)\b'
        ],
        'education': [
            r'\b(?:education|academic\s+background|educational\s+qualifications|academics|qualifications)\b'
        ],
        'projects': [
            r'\b(?:projects|key\s+projects|academic\s+projects|personal\s+projects)\b'
        ],
        'certifications': [
            r'\b(?:certifications|certificates|professional\s+certifications|licenses|courses)\b'
        ],
        'summary': [
            r'\b(?:professional\s+summary|summary|profile|career\s+objective|objective|about\s+me)\b'
        ]
    }

    DEGREE_PATTERNS = [
        r'\b(?:b\.?tech|b\.?e\.?|bachelor(?:\'s)?\s+of\s+technology|bachelor(?:\'s)?\s+of\s+engineering)\b',
        r'\b(?:m\.?tech|m\.?e\.?|master(?:\'s)?\s+of\s+technology|master(?:\'s)?\s+of\s+engineering)\b',
        r'\b(?:mca|master(?:\'s)?\s+of\s+computer\s+applications)\b',
        r'\b(?:bca|bachelor(?:\'s)?\s+of\s+computer\s+applications)\b',
        r'\b(?:b\.?sc|b\.?s\.?|bachelor(?:\'s)?\s+of\s+science)\b',
        r'\b(?:m\.?sc|m\.?s\.?|master(?:\'s)?\s+of\s+science)\b',
        r'\b(?:bachelor(?:\'s)?(?:\s+degree)?|master(?:\'s)?(?:\s+degree)?|ph\.?d|doctorate)\b',
        r'\b(?:diploma\s+in\s+[a-zA-Z\s]+)\b'
    ]

    def extract_text(self, file_path: str) -> str:
        """Extracts text from PDF or DOCX file."""
        if not os.path.exists(file_path):
            raise DocumentParsingError(f"File not found: {file_path}")

        ext = file_path.rsplit('.', 1)[-1].lower()

        if ext == 'pdf':
            return self._extract_from_pdf(file_path)
        elif ext == 'docx':
            return self._extract_from_docx(file_path)
        else:
            raise DocumentParsingError(f"Unsupported file format '.{ext}'. Only PDF and DOCX are supported.")

    def _extract_from_pdf(self, file_path: str) -> str:
        try:
            import pymupdf as fitz
            doc = fitz.open(file_path)
            text_pieces = []
            for page in doc:
                text_pieces.append(page.get_text())
            doc.close()
            full_text = "\n".join(text_pieces)
            return self.normalize_text(full_text)
        except Exception as e:
            raise DocumentParsingError(f"Error reading PDF file: {str(e)}")

    def _extract_from_docx(self, file_path: str) -> str:
        try:
            import docx
            doc = docx.Document(file_path)
            paragraphs = [p.text for p in doc.paragraphs if p.text.strip()]
            
            # Also extract tables if present
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
                        
            full_text = "\n".join(paragraphs)
            return self.normalize_text(full_text)
        except Exception as e:
            raise DocumentParsingError(f"Error reading DOCX file: {str(e)}")

    def normalize_text(self, text: str) -> str:
        """Removes extraneous characters and normalizes whitespace."""
        if not text:
            return ""
        # Replace non-breaking spaces and unusual whitespace
        text = text.replace('\xa0', ' ').replace('\r\n', '\n').replace('\r', '\n')
        # Standardize bullet symbols
        text = re.sub(r'[\u2022\u2023\u25E6\u2043\u2219]', '\n- ', text)
        # Remove consecutive blank lines
        text = re.sub(r'\n{3,}', '\n\n', text)
        return text.strip()

    def validate_content_sufficiency(self, text: str) -> Tuple[bool, str]:
        """
        Detects empty or scanned image documents where text extraction is insufficient.
        """
        cleaned = text.strip()
        if len(cleaned) < 60 or len(cleaned.split()) < 15:
            return False, "Insufficient text extracted (document may be empty, image-only, or scanned without an embedded text layer)."
        return True, "Sufficient text"

    def extract_metadata(self, text: str) -> Dict[str, Any]:
        """
        Extracts contact info, degree qualifications, estimated experience,
        and segmented sections.
        """
        return {
            'email': self._extract_email(text),
            'phone': self._extract_phone(text),
            'candidate_name': self._extract_name(text),
            'education': self._extract_education(text),
            'experience_years': self._estimate_experience_years(text),
            'sections': self._segment_sections(text)
        }

    def _extract_email(self, text: str) -> str | None:
        email_pattern = r'[a-zA-Z0-9_.+-]+@[a-zA-Z0-9-]+\.[a-zA-Z0-9-.]+'
        match = re.search(email_pattern, text)
        return match.group(0).lower() if match else None

    def _extract_phone(self, text: str) -> str | None:
        phone_pattern = r'(?:\+?\d{1,3}[-.\s]?)?\(?\d{3}\)?[-.\s]?\d{3}[-.\s]?\d{4}'
        match = re.search(phone_pattern, text)
        return match.group(0) if match else None

    def _extract_name(self, text: str) -> str:
        """Heuristic to guess candidate name from the top few lines."""
        lines = [line.strip() for line in text.split('\n') if line.strip()]
        for line in lines[:4]:
            # If line is short, doesn't contain email/urls, and looks like a name
            if 3 <= len(line) <= 40 and not re.search(r'[@:/\\0-9]', line):
                if not any(header in line.lower() for header in ['resume', 'curriculum', 'cv', 'summary', 'profile']):
                    return line
        return "Candidate"

    def _extract_education(self, text: str) -> List[str]:
        found_degrees = []
        for pat in self.DEGREE_PATTERNS:
            matches = re.findall(pat, text, re.IGNORECASE)
            for m in matches:
                normalized = m.strip().title()
                if normalized not in found_degrees:
                    found_degrees.append(normalized)
        return found_degrees

    def _estimate_experience_years(self, text: str) -> float:
        """
        Estimates years of experience from explicit statements or year ranges.
        """
        # Look for explicit patterns like '5+ years of experience', '3.5 years experience'
        exp_pattern = r'(\d+(?:\.\d+)?)\+?\s*(?:years?|yrs?)(?:\s+of)?\s+experience'
        matches = re.findall(exp_pattern, text, re.IGNORECASE)
        if matches:
            try:
                numbers = [float(m) for m in matches if float(m) <= 40]
                if numbers:
                    return max(numbers)
            except ValueError:
                pass

        # Check year date ranges like 2018 - 2023, 2019 - Present
        range_pattern = r'\b(20[0-2]\d|199\d)\s*[-–—to]+\s*(20[0-2]\d|Present|Current)\b'
        ranges = re.findall(range_pattern, text, re.IGNORECASE)
        total_years = 0.0
        current_year = 2026
        for start_str, end_str in ranges:
            try:
                start_yr = int(start_str)
                end_yr = current_year if end_str.lower() in ['present', 'current'] else int(end_str)
                if end_yr >= start_yr:
                    duration = end_yr - start_yr
                    total_years += min(duration, 15)  # Cap single span at 15 to avoid anomalies
            except ValueError:
                continue

        # If computed from spans, cap reasonably
        if total_years > 0:
            return round(min(total_years, 30.0), 1)

        return 0.0

    def _segment_sections(self, text: str) -> Dict[str, str]:
        """Segments resume into logical sections."""
        sections = {}
        lines = text.split('\n')
        current_section = 'summary'
        current_content = []

        for line in lines:
            trimmed = line.strip().lower()
            matched_header = None

            for sec_name, patterns in self.SECTION_PATTERNS.items():
                for pat in patterns:
                    if re.match(f'^{pat}[:]?$', trimmed, re.IGNORECASE):
                        matched_header = sec_name
                        break
                if matched_header:
                    break

            if matched_header:
                if current_content:
                    sections[current_section] = "\n".join(current_content).strip()
                current_section = matched_header
                current_content = []
            else:
                current_content.append(line)

        if current_content:
            sections[current_section] = "\n".join(current_content).strip()

        return sections
